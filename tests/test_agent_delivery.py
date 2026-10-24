import importlib
import json
from types import SimpleNamespace

import pytest
from docx import Document

from flight_agent.runtime.agent_contracts import AgentPlan, AgentScope, validate_bindings
from flight_agent.runtime.capabilities import CapabilityRegistry
from flight_agent.runtime.evidence import describe_evidence, verify_evidence
from flight_agent.runtime.observations import (
    CapabilityOutcome,
    Finding,
    Observation,
    OutputArtifact,
    ResultTable,
)
from flight_agent.runtime.requirements import TaskRequirements
from tests.agent_helpers import analysis_plan_data, observation_data, register_rmse
from tests.test_agent_evaluation import evaluation_plan


def adapter():
    return importlib.import_module("flight_agent.verification.adapters.delivery")


def context(tmp_path, outputs=None):
    return SimpleNamespace(
        directory=tmp_path,
        node_id="deliver",
        call_id="1" * 32,
        outputs=outputs
        or {
            "compute": CapabilityOutcome(
                observations=[Observation.model_validate(observation_data(conditions={}))]
            )
        },
    )


def requirements(formats, *, verification=False):
    data = evaluation_plan() if verification else analysis_plan_data()
    result = data["requirements"]
    result["delivery_formats"] = formats
    result["requested_outputs"].extend(
        {
            "id": f"report-{format}",
            "kind": "artifact",
            "node_id": "deliver",
            "item_id": format,
            "description": f"Deliver {format}.",
        }
        for format in formats
    )
    return TaskRequirements.model_validate(result)


def invoke(tmp_path, formats, *, outputs=None, verification=False):
    requested = requirements(formats, verification=verification)
    tools = CapabilityRegistry(agent_mode=True)
    tools.register(adapter().delivery_capability())
    result = tools.invoke(
        "deliver_results",
        {
            "title": "Generic measurement",
            "formats": formats,
            "requirements": requested.model_dump(mode="json"),
        },
        context(tmp_path, outputs),
        AgentScope(baseline_version=1, confirmed=True, allowed_effects=("write_delivery",)),
    )
    return CapabilityOutcome.model_validate(result), requested


def test_json_only_delivery_does_not_require_word_figures_or_matlab_artifacts(tmp_path):
    result, requested = invoke(tmp_path, ["json"])
    artifact = next(item for item in result.artifacts if item.id == "json")
    data = json.loads(verify_evidence(tmp_path, artifact).read_text())
    assert data["requirements"] == requested.model_dump(mode="json")
    assert data["evidence"]["observations"][0]["metric_id"] == "rmse"
    assert data["evidence"]["evaluation"] is None
    assert data["evaluation_requested"] is False
    assert not list(tmp_path.rglob("*.docx"))
    assert not list(tmp_path.rglob("*.png"))
    assert not list(tmp_path.rglob("*.mat"))


def test_word_delivery_preserves_new_metric_sources_conditions_invalid_reason_and_failed_rows(
    tmp_path,
):
    from tests.test_agent_evaluation import adapter as evaluation

    requested = requirements(["docx"], verification=True)
    observations = [
        Observation.model_validate(observation_data(value=3, conditions={})),
        Observation.model_validate(
            observation_data(
                metric_id="new-quality-metric",
                valid=False,
                value=None,
                reason="Window is not stationary.",
                conditions={"case": "transient"},
            )
        ),
    ]
    checked = evaluation().evaluate_criteria(
        requested.criteria, observations, overall_requested=True
    )
    outputs = {
        "compute": CapabilityOutcome(observations=observations),
        "evaluate": CapabilityOutcome(evaluation=checked),
    }
    result, _ = invoke(tmp_path, ["docx"], outputs=outputs, verification=True)
    artifact = next(item for item in result.artifacts if item.id == "docx")
    document = Document(verify_evidence(tmp_path, artifact))
    text = "\n".join(
        [paragraph.text for paragraph in document.paragraphs]
        + [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    )
    for expected in [
        "rmse",
        "new-quality-metric",
        "FAIL",
        "message-1",
        "Require synthetic",
        "transient",
        "Window is not stationary.",
        "Execution",
        "Validity",
    ]:
        assert expected in text


def test_delivery_checks_dependency_artifact_digests_before_writing_reports(tmp_path):
    (tmp_path / "evidence.json").write_text("{}")
    artifact = OutputArtifact(**describe_evidence(tmp_path, "evidence.json").model_dump())
    (tmp_path / "evidence.json").write_text("changed")
    with pytest.raises(ValueError, match="integrity"):
        invoke(tmp_path, ["json"], outputs={"compute": CapabilityOutcome(artifacts=[artifact])})
    assert not list(tmp_path.rglob("result.json"))


def test_explicitly_requested_word_failure_is_not_a_successful_json_fallback(tmp_path, monkeypatch):
    def fail(*args, **kwargs):
        raise OSError("Word destination failed")

    monkeypatch.setattr("docx.document.Document.save", fail)
    with pytest.raises(OSError, match="Word destination failed"):
        invoke(tmp_path, ["json", "docx"])
    assert not list(tmp_path.rglob("delivery-index.json"))


def delivery_plan():
    data = analysis_plan_data()
    requested = requirements(["json"])
    data["requirements"] = requested.model_dump(mode="json")
    data["allowed_effects"].append("write_delivery")
    data["nodes"].append(
        {
            "id": "deliver",
            "tool": "deliver_results",
            "tool_version": "1",
            "depends_on": ["compute"],
            "inputs": {
                "title": data["title"],
                "formats": ["json"],
                "requirements": requested.model_dump(mode="json"),
            },
        }
    )
    return data


@pytest.mark.parametrize("mutation", ["requirements", "formats", "producer", "format-id"])
def test_delivery_preflight_requires_exact_confirmed_requirements_and_producer_coverage(mutation):
    data = delivery_plan()
    if mutation == "requirements":
        data["nodes"][1]["inputs"]["requirements"]["objective"] = "Changed request"
    elif mutation == "formats":
        data["nodes"][1]["inputs"]["formats"] = ["docx"]
    elif mutation == "producer":
        data["nodes"][1]["depends_on"] = []
    else:
        data["requirements"]["requested_outputs"][1]["item_id"] = "unknown-report"
        data["nodes"][1]["inputs"]["requirements"] = data["requirements"]
    tools = CapabilityRegistry(agent_mode=True)
    register_rmse(tools, [])
    tools.register(adapter().delivery_capability())
    with pytest.raises(ValueError):
        validate_bindings(AgentPlan.model_validate(data), tools)


def test_delivery_has_explicit_write_permission_and_cannot_write_through_symlink(tmp_path):
    tools = CapabilityRegistry(agent_mode=True)
    tools.register(adapter().delivery_capability())
    inputs = {
        "title": "Generic",
        "formats": ["json"],
        "requirements": requirements(["json"]).model_dump(mode="json"),
    }
    with pytest.raises(ValueError, match="permission"):
        tools.invoke(
            "deliver_results",
            inputs,
            context(tmp_path),
            AgentScope(baseline_version=1, confirmed=True),
        )
    outside = tmp_path / "outside"
    outside.mkdir()
    (tmp_path / "delivery").symlink_to(outside, target_is_directory=True)
    with pytest.raises(ValueError, match="symlink"):
        tools.invoke(
            "deliver_results",
            inputs,
            context(tmp_path),
            AgentScope(baseline_version=1, confirmed=True, allowed_effects=("write_delivery",)),
        )
    assert list(outside.iterdir()) == []


def test_delivery_reconciliation_checks_original_identity_and_never_rewrites(tmp_path):
    tools = CapabilityRegistry(agent_mode=True)
    tools.register(adapter().delivery_capability())
    inputs = {
        "title": "Generic",
        "formats": ["json"],
        "requirements": requirements(["json"]).model_dump(mode="json"),
    }
    authorized = AgentScope(baseline_version=1, confirmed=True, allowed_effects=("write_delivery",))
    ctx = context(tmp_path)
    original = tools.invoke("deliver_results", inputs, ctx, authorized)
    before = {path: path.read_bytes() for path in tmp_path.rglob("*") if path.is_file()}
    assert tools.reconcile("deliver_results", inputs, ctx, authorized) == original
    assert {path: path.read_bytes() for path in before} == before
    with pytest.raises(ValueError, match="identity|input"):
        tools.reconcile("deliver_results", {**inputs, "title": "Reassigned"}, ctx, authorized)


def test_partial_delivery_discloses_missing_requested_outputs_instead_of_dropping_them(tmp_path):
    requested = requirements(["json"])
    data = requested.model_dump(mode="json")
    data["requested_outputs"].append(
        {
            "id": "missing-explanation",
            "kind": "finding",
            "node_id": "compute",
            "item_id": "explanation",
            "description": "Explain recorded evidence.",
        }
    )
    tools = CapabilityRegistry(agent_mode=True)
    tools.register(adapter().delivery_capability())
    result = tools.invoke(
        "deliver_results",
        {"title": "Partial evidence", "formats": ["json"], "requirements": data},
        context(tmp_path),
        AgentScope(baseline_version=1, confirmed=True, allowed_effects=("write_delivery",)),
    )
    artifact = next(item for item in result["artifacts"] if item["id"] == "json")
    payload = json.loads((tmp_path / artifact["path"]).read_text())
    assert payload["execution_status"] == "partial"
    assert payload["requirements"]["requested_outputs"][-1]["id"] == "missing-explanation"
    assert payload["unmet_requirements"][0]["requirement_id"] == "missing-explanation"


def test_word_only_delivery_preserves_producer_and_evidence_for_findings_and_tables(tmp_path):
    outputs = {
        node: CapabilityOutcome(
            findings=[
                Finding(
                    id="diagnosis",
                    node_id=node,
                    text="Measured finding.",
                    evidence=[f"{node}:finding-evidence"],
                )
            ],
            tables=[
                ResultTable(
                    id="samples",
                    node_id=node,
                    columns=["value"],
                    rows=[[1]],
                    evidence=[f"{node}:table-evidence"],
                )
            ],
        )
        for node in ["producer-a", "producer-b"]
    }
    result, _ = invoke(tmp_path, ["docx"], outputs=outputs)
    artifact = next(item for item in result.artifacts if item.id == "docx")
    document = Document(verify_evidence(tmp_path, artifact))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    for node in outputs:
        for expected in [
            f"{node}.diagnosis",
            f"{node}.samples",
            f"{node}:finding-evidence",
            f"{node}:table-evidence",
        ]:
            assert expected in text


def test_large_valid_table_delivery_can_be_reconciled_after_publication(tmp_path):
    tools = CapabilityRegistry(agent_mode=True)
    tools.register(adapter().delivery_capability())
    outputs = {
        "compute": CapabilityOutcome(
            tables=[
                ResultTable(
                    id="samples",
                    node_id="compute",
                    columns=[f"column-{index}" for index in range(40)],
                    rows=[[0] * 40 for _ in range(10_000)],
                )
            ]
        )
    }
    ctx = context(tmp_path, outputs)
    inputs = {
        "title": "Large but bounded table",
        "formats": ["json"],
        "requirements": requirements(["json"]).model_dump(mode="json"),
    }
    authorized = AgentScope(baseline_version=1, confirmed=True, allowed_effects=("write_delivery",))
    original = tools.invoke("deliver_results", inputs, ctx, authorized)
    assert tools.reconcile("deliver_results", inputs, ctx, authorized) == original


def test_index_serialized_byte_limit_is_checked_before_any_report_is_published(
    tmp_path, monkeypatch
):
    monkeypatch.setattr(adapter(), "MAX_INDEX_BYTES", 10, raising=False)
    with pytest.raises(ValueError, match="index.*size limit"):
        invoke(tmp_path, ["json", "docx"])
    assert list(tmp_path.iterdir()) == []
