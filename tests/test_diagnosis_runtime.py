import json
from datetime import UTC, datetime, timedelta
from pathlib import Path

import pytest
from docx import Document

from flight_agent.runtime.evidence import load_run
from flight_agent.verification.artifacts import verify_artifact
from flight_agent.verification.contracts import DiagnosisTask
from flight_agent.verification.diagnosis_contracts import (
    DiagnosisReview,
    DiagnosisTermination,
    EvidenceCitation,
)
from flight_agent.verification.runtime import VerificationRuntime
from tests.diagnosis_helpers import AdaptivePlanner, EvolvingPlanner, diagnostic_project
from tests.test_diagnosis_contracts import diagnosis_proposal
from tests.test_verification_runtime import ControlledBackend


def runtime_fixture(tmp_path, planner=None):
    planner = planner or AdaptivePlanner()
    backend = ControlledBackend()
    runtime = VerificationRuntime(tmp_path / "state", backend=backend, planner=planner)
    root = diagnostic_project(tmp_path / "source")
    return runtime, planner, backend, root


def propose(runtime, root):
    return runtime.submit(
        None, "Investigate the recorded pitch mismatch.", str(root), scenario="d2"
    )


def confirm(runtime, conversation):
    return runtime.confirm(conversation.id, conversation.task.baseline.version, trusted=True)


def test_d2_successful_evidence_drives_multiple_distinct_analysis_rounds_and_report(tmp_path):
    runtime, planner, backend, root = runtime_fixture(tmp_path)
    pending = propose(runtime, root)
    assert pending.status == "waiting_for_user", pending.error
    assert pending.task.scenario == "d2"
    assert not pending.task.analysis_runs
    assert not pending.task.baseline.confirmed
    completed = confirm(runtime, pending)
    assert completed.status == "completed", completed.error
    task = completed.task
    assert task.termination.reason == "evidence_sufficient"
    assert len(task.analysis_runs) == 2
    assert len(task.rounds) == len(task.plan_revisions) == len(planner.seen) == 2
    assert len({run.spec.fingerprint for run in task.analysis_runs}) == 2
    assert {run.spec.capability for run in task.analysis_runs} == {"compare_signals"}
    assert task.analysis_runs[0].spec.inputs["start"] == 0.0
    assert task.analysis_runs[1].spec.inputs["start"] == 2.0
    assert [item.status for item in task.assessments] == ["supported", "refuted"]
    assert task.termination.unknowns
    assert not backend.calls
    directory = Path(task.result_directory)
    for run in task.analysis_runs:
        assert load_run(directory, run.id, run.spec) == run
    for artifact in task.artifacts:
        verify_artifact(directory, artifact)
    report = next(item for item in task.artifacts if item.path.endswith(".docx"))
    text = "\n".join(paragraph.text for paragraph in Document(directory / report.path).paragraphs)
    assert all(run.id in text for run in task.analysis_runs)
    assert "synthetic" in text.lower()
    record = json.loads((directory / "delivery-v1/execution_record.json").read_text())
    assert record["status"] == "completed"
    assert record["termination"]["reason"] == "evidence_sufficient"
    assert len(record["analysis_runs"]) == 2
    assert all(call["status"] == "completed" for call in record["calls"])
    again = confirm(runtime, completed)
    assert again.task.analysis_runs == task.analysis_runs


def test_d2_changed_inputs_or_stale_approval_never_execute(tmp_path):
    runtime, planner, _, root = runtime_fixture(tmp_path)
    pending = propose(runtime, root)
    with pytest.raises(ValueError, match="version"):
        runtime.confirm(pending.id, 99, trusted=True)
    with pytest.raises(ValueError, match="consent|trusted"):
        runtime.confirm(pending.id, 1, trusted=False)
    path = root / "observed.json"
    data = json.loads(path.read_text())
    data["signals"]["pitch"]["values"][0] = 10.0
    path.write_text(json.dumps(data))
    rejected = confirm(runtime, pending)
    assert rejected.status == "waiting_for_user"
    assert not rejected.task.baseline.confirmed
    assert not rejected.task.analysis_runs
    assert "changed" in rejected.error.lower()


@pytest.mark.parametrize(
    "budget,reason", [("max_rounds", "round_budget"), ("max_calls", "call_budget")]
)
def test_d2_enforces_confirmed_budgets_without_inventing_a_conclusion(tmp_path, budget, reason):
    proposal = diagnosis_proposal()
    proposal["budget"] = {budget: 1}
    runtime, _, _, root = runtime_fixture(tmp_path, AdaptivePlanner(proposal))
    completed = confirm(runtime, propose(runtime, root))
    assert completed.status == "completed", completed.error
    assert completed.task.termination.reason == reason
    assert completed.task.termination.unknowns
    assert len(completed.task.analysis_runs) == 1
    assert completed.task.assessments[-1].status == "unknown"


def test_d2_repeated_specs_reuse_verified_receipts_then_stop_for_no_progress(tmp_path):
    runtime, planner, _, root = runtime_fixture(tmp_path, AdaptivePlanner(repeat=True))
    completed = confirm(runtime, propose(runtime, root))
    assert completed.status == "completed", completed.error
    assert completed.task.termination.reason == "no_progress"
    assert len(completed.task.analysis_runs) == 1
    assert len(completed.task.rounds) == 3
    assert sum(call.output.get("source") == "cache" for call in completed.task.calls) == 2
    assert len(planner.seen) == 3


def test_d2_unknown_citations_cannot_influence_next_execution(tmp_path):
    runtime, planner, _, root = runtime_fixture(
        tmp_path, AdaptivePlanner(invalid_citation=True)
    )
    failed = confirm(runtime, propose(runtime, root))
    assert failed.status == "failed"
    assert "evidence" in failed.error.lower() or "citation" in failed.error.lower()
    assert len(failed.task.analysis_runs) == 1
    assert failed.task.rounds == []
    assert failed.pending_decision.kind == "resume"
    planner.invalid_citation = False
    recovered = runtime.resume(failed.id)
    assert recovered.status == "completed", recovered.error
    assert len(recovered.task.analysis_runs) == 2
    assert len([call for call in recovered.task.calls if call.tool == "compare_signals"]) == 2


def test_d2_registry_limits_prevent_model_permission_escalation(tmp_path):
    proposal = diagnosis_proposal()
    proposal["allowed_capabilities"].append("write_candidate")
    runtime, _, _, root = runtime_fixture(tmp_path, AdaptivePlanner(proposal))
    rejected = propose(runtime, root)
    assert rejected.task is None
    assert rejected.error
    assert not list((tmp_path / "state").rglob("*.docx"))


def test_d2_receipt_before_acknowledgement_is_recovered_without_duplicate_dispatch(
    tmp_path, monkeypatch
):
    from flight_agent.verification import diagnosis

    runtime, _, _, root = runtime_fixture(tmp_path)
    original = diagnosis.persist_run
    writes = []

    def lost_ack(root, run):
        result = original(root, run)
        writes.append(run.id)
        if len(writes) == 1:
            raise OSError("Controlled lost receipt acknowledgement")
        return result

    monkeypatch.setattr(diagnosis, "persist_run", lost_ack)
    pending = propose(runtime, root)
    failed = confirm(runtime, pending)
    assert failed.status == "failed"
    assert failed.task.pending_analysis is not None
    run_id = failed.task.pending_analysis.id
    restored = VerificationRuntime(tmp_path / "state", planner=runtime.planner)
    completed = restored.resume(failed.id)
    assert completed.status == "completed", completed.error
    assert completed.task.analysis_runs[0].id == run_id
    assert writes.count(run_id) == 1
    assert len(completed.task.analysis_runs) == 2


def test_d2_unknown_inflight_execution_stops_instead_of_redispatching(tmp_path, monkeypatch):
    runtime, _, _, root = runtime_fixture(tmp_path)
    pending = propose(runtime, root)
    calls = []

    def interrupted(*args, **kwargs):
        calls.append(True)
        raise KeyboardInterrupt("Controlled interruption before a receipt")

    monkeypatch.setattr(runtime.diagnosis.tools, "invoke", interrupted)
    with pytest.raises(KeyboardInterrupt):
        confirm(runtime, pending)
    restored = VerificationRuntime(tmp_path / "state", planner=runtime.planner)
    completed = restored.resume(pending.id)
    assert completed.status == "completed", completed.error
    assert completed.task.termination.reason == "unknown_execution"
    assert not completed.task.analysis_runs
    assert calls == [True]


def test_d2_switch_to_d1_preserves_task_and_pauses_without_reusing_data_root(tmp_path):
    runtime, _, _, root = runtime_fixture(tmp_path)
    pending = propose(runtime, root)
    paused = runtime.submit(pending.id, "Switch to verification.", scenario="d1")
    assert paused.scenario == "d1"
    assert paused.project_root is None
    assert paused.output_root is None
    assert paused.task == pending.task
    assert paused.archived_tasks == []
    assert paused.pending_decision.kind == "clarify"
    with pytest.raises(ValueError, match="workflow|stage"):
        confirm(runtime, paused)


def test_d2_expired_wall_clock_budget_survives_restart_without_new_dispatch(tmp_path):
    runtime, planner, _, root = runtime_fixture(
        tmp_path, AdaptivePlanner(invalid_citation=True)
    )
    failed = confirm(runtime, propose(runtime, root))
    failed.task.started_at = (
        datetime.now(UTC) - timedelta(seconds=failed.task.baseline.budget.max_seconds + 1)
    ).isoformat()
    runtime.store.save(failed)
    planner.invalid_citation = False
    restored = VerificationRuntime(tmp_path / "state", planner=planner)
    stopped = restored.resume(failed.id)
    assert stopped.status == "completed", stopped.error
    assert stopped.task.termination.reason == "time_budget"
    assert stopped.task.dispatch_count == 1
    assert len(stopped.task.analysis_runs) == 1


def test_d2_invalid_numerical_evidence_is_reported_as_unknown_not_a_supported_cause(tmp_path):
    class MissingEvidencePlanner(AdaptivePlanner):
        def review_diagnosis(self, task, *, tools):
            return DiagnosisReview.model_validate({
                "action": "stop", "stop_reason": "missing_input",
                "reason": "The recorded signals have incompatible units.",
                "assessments": [
                    {
                        "hypothesis_id": hypothesis.id, "status": "unknown",
                        "reason": "Comparable, correctly scaled observations are missing.",
                    }
                    for hypothesis in task.baseline.hypotheses
                ],
                "unknowns": ["No automatic unit conversion or causal assumption is permitted."],
                "summary": {
                    "problem": "The recorded signals have incompatible units.",
                    "localization": "Comparable pitch observations are missing.",
                    "responsibility": "undetermined",
                    "reason": "Invalid numerical evidence cannot support responsibility.",
                    "recommendations": [
                        "Provide independently verified, consistently scaled data.",
                    ],
                },
            })

    runtime, _, _, root = runtime_fixture(tmp_path, MissingEvidencePlanner())
    path = root / "reference.json"
    payload = json.loads(path.read_text())
    payload["signals"]["pitch"]["unit"] = "rad"
    path.write_text(json.dumps(payload))
    stopped = confirm(runtime, propose(runtime, root))
    assert stopped.status == "completed", stopped.error
    assert stopped.task.analysis_runs[0].output["valid"] is False
    assert stopped.task.termination.reason == "missing_input"
    assert {item.status for item in stopped.task.assessments} == {"unknown"}


def test_d2_resume_uses_immutable_snapshot_not_later_source_changes(tmp_path):
    runtime, planner, _, root = runtime_fixture(
        tmp_path, AdaptivePlanner(invalid_citation=True)
    )
    failed = confirm(runtime, propose(runtime, root))
    (root / "observed.json").write_text('{"not":"the confirmed input"}')
    planner.invalid_citation = False
    done = runtime.resume(failed.id)
    assert done.status == "completed", done.error
    assert done.task.analysis_runs[-1].output["metrics"]["mean_error"]["value"] == 1.0


def test_d2_resume_rejects_tampered_snapshot_metadata_before_further_execution(tmp_path):
    runtime, planner, _, root = runtime_fixture(
        tmp_path, AdaptivePlanner(invalid_citation=True)
    )
    failed = confirm(runtime, propose(runtime, root))
    snapshot = failed.task.input_snapshot
    failed.task.input_snapshot = snapshot.model_copy(update={
        "manifest": snapshot.manifest.model_copy(update={"conditions": {"tampered": True}})
    })
    runtime.store.save(failed)
    planner.invalid_citation = False
    stopped = runtime.resume(failed.id)
    assert stopped.status == "failed"
    assert "integrity" in stopped.error.lower()
    assert stopped.task.dispatch_count == 1


def test_d2_report_rechecks_citations_instead_of_trusting_mutable_task_projection(tmp_path):
    from flight_agent.verification.diagnosis_reports import deliver_diagnosis_report

    runtime, _, _, root = runtime_fixture(tmp_path)
    done = confirm(runtime, propose(runtime, root))
    task = done.task
    task.assessments[0] = task.assessments[0].model_copy(
        update={"evidence": [EvidenceCitation(run_id="f" * 32, metric="mean_error")]}
    )
    with pytest.raises(ValueError, match="evidence|citation"):
        deliver_diagnosis_report(task, done.messages)


@pytest.mark.parametrize("failed", [False, True])
def test_d2_followup_clarification_preserves_original_task_lifecycle(tmp_path, failed):
    runtime, planner, _, root = runtime_fixture(
        tmp_path, AdaptivePlanner(invalid_citation=failed)
    )
    previous = confirm(runtime, propose(runtime, root))
    planner.initial = {"action": "clarify", "question": "Which recorded signal should be compared?"}
    clarified = runtime.submit(previous.id, "Explain a different observation.")
    assert clarified.status == previous.status
    assert clarified.error == previous.error
    assert clarified.task.id == previous.task.id
    if failed:
        assert clarified.pending_decision.kind == "resume"


def test_d2_revised_hypotheses_drive_the_next_plan_and_structured_deliverables(tmp_path):
    runtime, _, backend, root = runtime_fixture(tmp_path, EvolvingPlanner())
    pending = propose(runtime, root)
    original_hypotheses = pending.task.baseline.hypotheses
    done = confirm(runtime, pending)
    assert done.status == "completed", done.error
    task = done.task
    assert task.baseline.hypotheses == original_hypotheses
    assert [(item.id, item.priority) for item in task.current_hypotheses] == [
        ("bias", "high"), ("transient", "medium"), ("alignment", "low"),
    ]
    assert task.plan_revisions[1].analyses[0].hypothesis_ids == [
        "bias", "transient", "alignment",
    ]
    assert len(task.rounds) == len(task.analysis_runs) == 2
    assert task.rounds[0].review.hypothesis_revision is not None
    assert task.rounds[1].review.hypothesis_revision is None
    assert task.termination.reason == "missing_input"
    assert task.rounds[-1].review.summary.responsibility == "undetermined"
    assert not backend.calls
    directory = Path(task.result_directory)
    payload = json.loads((directory / "delivery-v1/diagnosis.json").read_text())
    assert [item["id"] for item in payload["current_hypotheses"]] == [
        "bias", "transient", "alignment",
    ]
    assert payload["summary"]["responsibility"] == "undetermined"
    assert payload["hypothesis_revisions"][0]["plan_version"] == 1
    assert payload["hypothesis_revisions"][0]["revision"]["evidence"][0]["run_id"] == (
        task.analysis_runs[0].id
    )
    report = Document(directory / "delivery-v1/diagnosis.docx")
    text = "\n".join(item.text for item in report.paragraphs)
    assert "Problem localization" in text
    assert "undetermined" in text
    assert "Verify acquisition timing and signal scaling independently." in text
    assert "HIGH" in text and "alignment" in text


def test_d2_resume_preserves_revisions_and_reuses_existing_evidence(tmp_path):
    planner = EvolvingPlanner(fail_second=True)
    runtime, _, _, root = runtime_fixture(tmp_path, planner)
    failed = confirm(runtime, propose(runtime, root))
    assert failed.status == "failed"
    assert len(failed.task.rounds) == 1
    assert len(failed.task.analysis_runs) == 2
    original_runs = list(failed.task.analysis_runs)
    first_revision = failed.task.rounds[0]
    planner.fail_second = False
    restarted = VerificationRuntime(tmp_path / "state", planner=planner)
    done = restarted.resume(failed.id)
    assert done.status == "completed", done.error
    assert done.task.analysis_runs == original_runs
    assert done.task.dispatch_count == 2
    assert done.task.rounds[0] == first_revision
    assert done.task.current_hypotheses[-1].id == "alignment"


@pytest.mark.parametrize("fault", ["drop", "citation", "scope", "attribution"])
def test_d2_invalid_revision_stops_before_dispatching_the_next_analysis(tmp_path, fault):
    runtime, _, _, root = runtime_fixture(tmp_path, EvolvingPlanner(revision_fault=fault))
    failed = confirm(runtime, propose(runtime, root))
    assert failed.status == "failed"
    assert failed.task.dispatch_count == 1
    assert len(failed.task.plan_revisions) == 1
    assert failed.task.rounds == []
    assert failed.task.current_hypotheses == failed.task.baseline.hypotheses


def test_d2_new_reviews_require_explicit_structured_findings(tmp_path):
    class NoSummaryPlanner(AdaptivePlanner):
        def review_diagnosis(self, task, *, tools):
            return super().review_diagnosis(task, tools=tools).model_copy(
                update={"summary": None}
            )

    runtime, _, _, root = runtime_fixture(tmp_path, NoSummaryPlanner())
    failed = confirm(runtime, propose(runtime, root))
    assert failed.status == "failed"
    assert "summary" in failed.error.lower()
    assert failed.task.dispatch_count == 1
    assert not failed.task.rounds


@pytest.mark.parametrize("fault", ["future_evidence", "revision", "round_evidence"])
def test_d2_history_replay_rejects_changed_revision_or_noncausal_evidence_order(tmp_path, fault):
    from flight_agent.verification.diagnosis_reports import deliver_diagnosis_report

    planner = EvolvingPlanner(fail_second=True)
    runtime, _, _, root = runtime_fixture(tmp_path, planner)
    failed = confirm(runtime, propose(runtime, root))
    task = failed.task
    assert len(task.rounds) == 1, failed.error
    first = task.rounds[0]
    revision = first.review.hypothesis_revision
    if fault == "future_evidence":
        revision = revision.model_copy(update={
            "evidence": [EvidenceCitation(run_id=task.analysis_runs[1].id, metric="mean_error")],
        })
    elif fault == "revision":
        revision = revision.model_copy(update={"hypotheses": revision.hypotheses[:-1]})
    else:
        first = first.model_copy(update={"run_ids": [task.analysis_runs[1].id]})
    task.rounds[0] = first.model_copy(update={
        "review": first.review.model_copy(update={"hypothesis_revision": revision}),
    })
    runtime.store.save(failed)
    planner.fail_second = False
    stopped = runtime.resume(failed.id)
    assert stopped.status == "failed"
    assert stopped.task.dispatch_count == 2
    assert len(stopped.task.rounds) == 1
    task.termination = task.termination or DiagnosisTermination(
        reason="missing_input", detail="Cannot trust this altered history.",
        created_at="2026-10-20T00:00:02Z", unknowns=["History is inconsistent."],
    )
    with pytest.raises(ValueError, match="evidence|hypothes|round|plan"):
        deliver_diagnosis_report(task, failed.messages)


def test_d2_legacy_records_report_missing_summary_without_inventing_one(tmp_path):
    from flight_agent.verification.diagnosis_reports import deliver_diagnosis_report

    runtime, _, _, root = runtime_fixture(tmp_path)
    done = confirm(runtime, propose(runtime, root))
    payload = done.task.model_dump()
    for item in payload["baseline"]["hypotheses"]:
        item.pop("priority")
    for item in payload["rounds"]:
        item["review"].pop("summary")
        item["review"].pop("hypothesis_revision")
    legacy = DiagnosisTask.model_validate(payload)
    assert legacy.current_hypotheses == legacy.baseline.hypotheses
    deliver_diagnosis_report(legacy, done.messages)
    directory = Path(legacy.result_directory) / "delivery-v1"
    summary = json.loads((directory / "diagnosis.json").read_text())["summary"]
    assert summary is None
    text = "\n".join(item.text for item in Document(directory / "diagnosis.docx").paragraphs)
    assert "Structured diagnosis unavailable" in text
