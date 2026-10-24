import hashlib
import json
from io import BytesIO
from pathlib import Path
from typing import Annotated, Literal, Protocol, Self

from docx import Document
from pydantic import Field, JsonValue, TypeAdapter, model_validator

from ...runtime.agent_contracts import AgentNode, AgentPlan
from ...runtime.capabilities import Capability
from ...runtime.capability_metadata import CapabilityMetadata
from ...runtime.contracts import Digest, Name, RunId, fingerprint
from ...runtime.evidence import describe_evidence, publish_evidence, verify_evidence
from ...runtime.observations import (
    CapabilityOutcome,
    OutputArtifact,
    TaskResult,
    combine_outcomes,
)
from ...runtime.requirements import CapabilityGap, TaskRequirements
from ...runtime.resources import AgentContract
from ..intake import safe_directory
from .evaluation import OutcomeContext

Format = Literal["json", "docx"]
_CALL_ID = TypeAdapter(RunId)
MAX_INDEX_BYTES = 4_000_000


class DeliveryContext(OutcomeContext, Protocol):
    @property
    def directory(self) -> Path: ...

    @property
    def node_id(self) -> str: ...

    @property
    def call_id(self) -> str: ...


class DeliveryInput(AgentContract):
    title: Annotated[str, Field(min_length=1, max_length=200)]
    formats: Annotated[list[Format], Field(min_length=1, max_length=2)]
    requirements: TaskRequirements

    @model_validator(mode="after")
    def requested_formats(self) -> Self:
        if self.formats != self.requirements.delivery_formats:
            raise ValueError("Delivery formats differ from the confirmed requirements")
        return self


class _DeliveryIndex(AgentContract):
    input_sha256: Digest
    dependency_output_sha256: dict[Name, Digest]
    outcome: CapabilityOutcome
    result: TaskResult


def _prefix(context: DeliveryContext) -> str:
    safe_directory(context.directory)
    return f"delivery/{_CALL_ID.validate_python(context.call_id)}"


def _dependencies(context: DeliveryContext) -> dict[str, CapabilityOutcome]:
    outputs = {
        name: CapabilityOutcome.model_validate(value) for name, value in context.outputs.items()
    }
    for output in outputs.values():
        for artifact in output.artifacts:
            verify_evidence(context.directory, artifact)
    return outputs


def _write_word(inputs: DeliveryInput, evidence: CapabilityOutcome, result: TaskResult) -> bytes:
    document = Document()
    document.add_heading(inputs.title, 0)
    document.add_heading("Execution", 1)
    document.add_paragraph(
        f"Evidence processing: {result.execution_status}. "
        "Execution, numerical validity and acceptance are separate results. "
        "This report is not flight qualification."
    )
    document.add_paragraph(inputs.requirements.objective)
    document.add_heading("Observations and Validity", 1)
    for item in evidence.observations:
        document.add_paragraph(
            f"{item.node_id}.{item.metric_id} ({item.label}): "
            f"{item.value if item.value is not None else item.special or 'unavailable'} "
            f"{item.unit}; "
            f"Validity: {item.valid}; {item.reason}; "
            f"conditions: {json.dumps(item.conditions, ensure_ascii=False, sort_keys=True)}; "
            f"evidence: {', '.join(item.evidence)}"
        )
    document.add_heading("Criteria and sources", 1)
    for criterion in inputs.requirements.criteria:
        document.add_paragraph(
            f"{criterion.id}: {criterion.metric.node_id}.{criterion.metric.metric_id} "
            f"{criterion.operator} {criterion.threshold} {criterion.unit}; "
            f"conditions: {json.dumps(criterion.conditions, ensure_ascii=False, sort_keys=True)}"
        )
        document.add_paragraph(
            f"Source {criterion.source.kind}: {criterion.source.reference}; "
            f"{criterion.source.quote}; version: {criterion.source.version or 'not provided'}; "
            f"SHA-256: {criterion.source.sha256 or 'message reference'}"
        )
    document.add_heading("Evaluation", 1)
    if evidence.evaluation is None:
        document.add_paragraph(
            "Not requested." if inputs.requirements.mode == "analysis" else "Not available."
        )
    else:
        document.add_paragraph(
            f"Overall: {evidence.evaluation.status.upper()}"
            if evidence.evaluation.status
            else "No overall acceptance decision was requested."
        )
        for check in evidence.evaluation.checks:
            document.add_paragraph(
                f"{check.criterion_id}: {check.status.upper()}; "
                f"{check.value if check.value is not None else check.special or 'unavailable'} "
                f"{check.value_unit or ''} {check.operator} {check.threshold} {check.unit}; "
                f"{check.reason}; evidence: {', '.join(check.evidence)}"
            )
    for finding in evidence.findings:
        document.add_paragraph(
            f"{finding.node_id or 'unrecorded-producer'}.{finding.id}: "
            f"{finding.status}; {finding.text}; evidence: {', '.join(finding.evidence)}"
        )
    for table in evidence.tables:
        document.add_heading(f"{table.node_id or 'unrecorded-producer'}.{table.id}", 2)
        document.add_paragraph(f"Evidence: {', '.join(table.evidence)}")
        rendered = document.add_table(rows=1, cols=len(table.columns))
        for cell, label in zip(rendered.rows[0].cells, table.columns, strict=True):
            cell.text = label
        for row in table.rows:
            for cell, value in zip(rendered.add_row().cells, row, strict=True):
                cell.text = json.dumps(value, ensure_ascii=False)
    for gap in result.unmet_requirements:
        document.add_paragraph(f"Unmet {gap.requirement_id}: {gap.limitation}; {gap.next_step}")
    for artifact in evidence.artifacts:
        document.add_paragraph(
            f"Evidence file: {artifact.path}; SHA-256: {artifact.sha256}; bytes: {artifact.size}"
        )
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def _dependency_hashes(outputs: dict[str, CapabilityOutcome]) -> dict[str, str]:
    return {name: fingerprint(value.model_dump(mode="json")) for name, value in outputs.items()}


def _index_outcome(context: DeliveryContext, index: _DeliveryIndex) -> CapabilityOutcome:
    name = f"{_prefix(context)}/delivery-index.json"
    artifact = OutputArtifact(
        **describe_evidence(context.directory, name).model_dump(),
        id="delivery-index",
        node_id=context.node_id,
    )
    return CapabilityOutcome(artifacts=[*index.outcome.artifacts, artifact])


def _reconcile(context: DeliveryContext, inputs: DeliveryInput) -> CapabilityOutcome | None:
    name = f"{_prefix(context)}/delivery-index.json"
    path = context.directory / name
    safe_directory(path.parent)
    if not path.exists():
        return None
    descriptor = describe_evidence(context.directory, name)
    if descriptor.size > MAX_INDEX_BYTES:
        raise ValueError("Delivery index exceeds its size limit")
    index = _DeliveryIndex.model_validate_json(path.read_bytes())
    outputs = _dependencies(context)
    if index.input_sha256 != fingerprint(inputs.model_dump(mode="json")) or (
        index.dependency_output_sha256 != _dependency_hashes(outputs)
    ):
        raise ValueError("Delivery index input or dependency identity mismatch")
    if {artifact.id for artifact in index.outcome.artifacts} != set(inputs.formats):
        raise ValueError("Delivery index output identity mismatch")
    for artifact in index.outcome.artifacts:
        if artifact.node_id != context.node_id or artifact.path != (
            f"{_prefix(context)}/result.{artifact.id}"
        ):
            raise ValueError("Delivery artifact identity mismatch")
        verify_evidence(context.directory, artifact)
    index.result.validate_requirements(inputs.requirements)
    return _index_outcome(context, index)


def _deliver(context: DeliveryContext, inputs: DeliveryInput) -> CapabilityOutcome:
    existing = _reconcile(context, inputs)
    if existing is not None:
        return existing
    outputs = _dependencies(context)
    evidence = combine_outcomes(outputs)
    gaps = [
        CapabilityGap(
            requirement_id=output.id,
            limitation="Requested output is absent from the verified dependencies.",
            next_step="Supply the missing evidence through a newly confirmed plan.",
        )
        for output in evidence.missing_outputs(inputs.requirements)
        if not (
            output.kind == "artifact"
            and output.node_id == context.node_id
            and output.item_id in inputs.formats
        )
    ]
    result = TaskResult(
        **evidence.model_dump(mode="json"),
        execution_status="partial" if gaps else "completed",
        unmet_requirements=gaps,
    )
    payload: dict[str, JsonValue] = {
        "title": inputs.title,
        "requirements": inputs.requirements.model_dump(mode="json"),
        "evidence": evidence.model_dump(mode="json"),
        "execution_status": result.execution_status,
        "evaluation_requested": inputs.requirements.mode == "verification",
        "unmet_requirements": [gap.model_dump(mode="json") for gap in gaps],
    }
    prepared = []
    for format in inputs.formats:
        content = (
            (json.dumps(payload, ensure_ascii=False, indent=2, allow_nan=False) + "\n").encode(
                "utf-8"
            )
            if format == "json"
            else _write_word(inputs, evidence, result)
        )
        name = f"{_prefix(context)}/result.{format}"
        artifact = OutputArtifact(
            path=name,
            sha256=hashlib.sha256(content).hexdigest(),
            size=len(content),
            id=format,
            node_id=context.node_id,
        )
        prepared.append((artifact, content))
    artifacts = [artifact for artifact, _ in prepared]
    result = TaskResult.model_validate(
        {**result.model_dump(mode="json"), "artifacts": [*evidence.artifacts, *artifacts]}
    )
    result.validate_requirements(inputs.requirements)
    index = _DeliveryIndex(
        input_sha256=fingerprint(inputs.model_dump(mode="json")),
        dependency_output_sha256=_dependency_hashes(outputs),
        outcome=CapabilityOutcome(artifacts=artifacts),
        result=result,
    )
    index_bytes = index.model_dump_json().encode("utf-8")
    if len(index_bytes) > MAX_INDEX_BYTES:
        raise ValueError("Delivery index exceeds its size limit")
    for artifact, content in prepared:
        publish_evidence(context.directory, artifact.path, content)
        verify_evidence(context.directory, artifact)
    publish_evidence(context.directory, f"{_prefix(context)}/delivery-index.json", index_bytes)
    return _index_outcome(context, index)


def _validate_plan(plan: AgentPlan, node: AgentNode) -> None:
    if node.bindings:
        raise ValueError("Delivery requirements must be literal confirmed inputs")
    inputs = DeliveryInput.model_validate(node.inputs)
    if inputs.requirements != plan.requirements:
        raise ValueError("Delivery inputs differ from the confirmed requirements")
    gap_ids = {gap.requirement_id for gap in plan.gaps}
    for output in plan.requirements.requested_outputs:
        if output.id in gap_ids:
            continue
        if output.node_id == node.id:
            if output.kind != "artifact" or output.item_id not in inputs.formats:
                raise ValueError("Requested delivery artifact is not a supported format")
        elif output.node_id not in node.depends_on:
            raise ValueError("Delivery requires directly declared output producers")
    if plan.evaluation_node is not None and plan.evaluation_node not in node.depends_on:
        raise ValueError("Delivery requires the confirmed evaluation output")


def delivery_capability() -> Capability[DeliveryContext, DeliveryInput, CapabilityOutcome]:
    return Capability(
        name="deliver_results",
        version="1",
        description="Deliver confirmed generic evidence and optional evaluation as JSON or Word.",
        input_model=DeliveryInput,
        output_model=CapabilityOutcome,
        effect="write_delivery",
        operation=_deliver,
        reconcile=_reconcile,
        validate_plan=_validate_plan,
        metadata=CapabilityMetadata(
            validity_policy={
                "version": "1",
                "evidence": "Verified dependency receipts and file digests.",
                "delivery": "Immutable per-call files; no unrequested acceptance decision.",
            },
            planning_notes=(
                "Formats json/docx only, logical artifact IDs equal those formats. Pass the exact "
                "confirmed requirements; directly depend on every output producer and evaluation. "
                "Dependency receipts must be verified by the executor before building context."
            ),
        ),
    )
