import json
from pathlib import Path
from typing import Literal

from docx import Document
from docx.shared import Inches
from pydantic import JsonValue, field_validator

from ..models import Contract, Nonempty
from ..runtime.contracts import fingerprint
from ..session import _write_json
from ..upload_models import finite_json
from .artifacts import describe_artifact, verify_artifact
from .contracts import (
    AnalysisKind,
    AnalysisResult,
    Artifact,
    Evaluation,
    EvaluationBaseline,
    Identifier,
    ProjectDescriptor,
)
from .evaluation import evaluate

REPORT_FORMAT_VERSION = "1"


class ReportRequest(Contract):
    task_id: Identifier
    title: Nonempty
    project: ProjectDescriptor
    baseline: EvaluationBaseline
    result: AnalysisResult
    evaluation: Evaluation
    derived_artifacts: list[Artifact] = []
    grouping: Literal["combined", "separate"] = "combined"
    explanation: str | None = None
    explanation_evidence: list[str] = []
    explanation_error: str | None = None
    execution_record: dict[str, JsonValue] = {}

    @field_validator("execution_record")
    @classmethod
    def finite_record(cls, value: dict[str, JsonValue]) -> dict[str, JsonValue]:
        finite_json(value)
        return value


def _write_report(
    request: ReportRequest, directory: Path, report_name: str, kinds: list[AnalysisKind]
) -> None:
    document = Document()
    document.add_heading(request.title, 0)
    document.add_paragraph(
        f"Task {request.task_id}; confirmed baseline v{request.baseline.version}; "
        f"trim point {request.baseline.trim_point}"
    )
    statuses = [request.evaluation.by_analysis[kind] for kind in kinds]
    status = (
        "undetermined" if "undetermined" in statuses else "fail" if "fail" in statuses else "pass"
    )
    document.add_heading(f"Verification: {status.upper()}", 1)
    document.add_paragraph(
        "This is a result against the confirmed criteria, not full-envelope flight qualification. "
        "Independent reference results and agreed tolerances are required "
        "for engineering acceptance."
    )
    document.add_heading("Project and analysis conditions", 1)
    config = request.project.config
    for text in (
        f"Project: {config.name}",
        f"Project fingerprint: {request.project.fingerprint}",
        f"Backend: {config.backend}; MATLAB: {request.result.matlab_version}",
        f"Model: {config.model}; initializer: {config.initializer}",
        f"Input: {config.input_signal.name} [{config.input_signal.unit}]; "
        f"response: {config.response_signal.name} [{config.response_signal.unit}]",
        f"Operating point: {config.operating_point_mode}; feedback sign: {config.feedback_sign}",
        "Step settling band: 5% of the response change; measured from the detected step time.",
        "Steady-state screening: " + request.baseline.steady_state.model_dump_json(),
    ):
        document.add_paragraph(text)
    manifest = {item.path: item for item in request.result.artifacts}
    for kind in kinds:
        document.add_heading(kind, 1)
        record = request.result.analysis(kind)
        document.add_paragraph(
            f"Execution: {record.execution_status if record else 'missing'}; "
            f"criterion result: {request.evaluation.by_analysis[kind].upper()}"
        )
        table = document.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        for cell, text in zip(
            table.rows[0].cells, ("Metric", "Value", "Criterion", "Result"), strict=True
        ):
            cell.text = text
        for check in request.evaluation.checks:
            if check.analysis != kind:
                continue
            value = (
                f"{check.value:g}" if check.value is not None else check.special or "unavailable"
            )
            values = (
                check.metric,
                f"{value} {check.value_unit or check.unit}",
                f"{check.operator} {check.threshold:g} {check.unit}",
                check.status.upper(),
            )
            for cell, text in zip(table.add_row().cells, values, strict=True):
                cell.text = text
            document.add_paragraph(f"{check.metric}: {check.reason}")
            if check.evidence:
                document.add_paragraph("Evidence: " + ", ".join(check.evidence))
        if record is not None:
            for diagnostic in record.diagnostics:
                document.add_paragraph("Diagnostic: " + diagnostic)
            for name in record.artifact_paths:
                item = manifest[name]
                if item.role == "figure":
                    document.add_picture(str(verify_artifact(directory, item)), width=Inches(6))
                document.add_paragraph(f"Evidence file: {item.path}; SHA-256: {item.sha256}")
    document.add_heading("LLM explanation (not a numerical authority)", 1)
    if request.explanation:
        document.add_paragraph(request.explanation)
        document.add_paragraph("Referenced evidence: " + ", ".join(request.explanation_evidence))
    else:
        document.add_paragraph(
            "No LLM explanation was produced. "
            + (request.explanation_error or "Only deterministic evidence is reported.")
        )
    for warning in request.evaluation.warnings:
        document.add_paragraph("Limitation: " + warning)
    document.add_heading("Reproduction and evidence", 1)
    document.add_paragraph(
        "Structured inputs, results and execution history accompany this report."
    )
    document.add_paragraph(
        json.dumps(request.result.provenance, ensure_ascii=False, sort_keys=True)
    )
    path = directory / report_name
    temporary = path.with_suffix(".tmp.docx")
    document.save(str(temporary))
    temporary.replace(path)


def deliver_report(request: ReportRequest, directory: Path) -> list[Artifact]:
    if not request.baseline.confirmed:
        raise ValueError("A confirmed evaluation baseline is required before delivery")
    if (
        request.project.fingerprint != request.baseline.project_fingerprint
        or request.project.fingerprint != request.result.project_fingerprint
    ):
        raise ValueError("Report project context differs from its evidence")
    expected = evaluate(request.baseline, request.result)
    if request.evaluation != expected:
        raise ValueError("Report evaluation must equal the deterministic evidence-based result")
    known = {item.path for item in request.result.artifacts}
    for record in request.result.analyses:
        if not set(record.artifact_paths).issubset(known):
            raise ValueError(f"Unregistered analysis evidence: {record.kind}")
        if record.valid and not (
            any(name.endswith(".png") for name in record.artifact_paths)
            and any(name.endswith(".mat") for name in record.artifact_paths)
        ):
            raise ValueError(f"Valid analysis lacks figure or raw evidence: {record.kind}")
    for item in request.result.artifacts:
        verify_artifact(directory, item)
    for item in request.derived_artifacts:
        if item.role != "record" or not item.path.startswith("evaluation/"):
            raise ValueError("Only registered evaluation receipts are derived evidence")
        verify_artifact(directory, item)
    prefix = f"delivery-v{request.baseline.version}"
    target = directory / prefix
    if target.is_symlink():
        raise ValueError("Unsafe report directory")
    target.mkdir(parents=True, exist_ok=True)
    coverage: dict[str, list[AnalysisKind]]
    if request.grouping == "combined":
        coverage = {f"{prefix}/verification.docx": request.baseline.analyses}
    else:
        coverage = {
            f"{prefix}/{kind.replace('_', '-')}.docx": [kind] for kind in request.baseline.analyses
        }
    render_key = fingerprint(
        {
            "format_version": REPORT_FORMAT_VERSION,
            **request.model_dump(mode="json", exclude={"execution_record"}),
        }
    )
    reused = False
    index_name = f"{prefix}/artifact-index.json"
    if (directory / index_name).exists() or (directory / index_name).is_symlink():
        index_file = describe_artifact(directory, index_name)
        index = json.loads(verify_artifact(directory, index_file).read_text("utf-8"))
        if not isinstance(index, dict):
            raise ValueError("Invalid report artifact index")
        if index.get("report_fingerprint") == render_key:
            if index.get("reports") != coverage:
                raise ValueError("Cached report coverage does not match the requested delivery")
            cached = {
                item.path: item
                for item in (Artifact.model_validate(value) for value in index["artifacts"])
            }
            for name in coverage:
                if name not in cached or cached[name].role != "report":
                    raise ValueError("Cached report is missing its registered artifact")
                verify_artifact(directory, cached[name])
            reused = True
    if not reused:
        for name, kinds in coverage.items():
            _write_report(request, directory, name, kinds)
    _write_json(
        target / "evaluation.json",
        {
            "task_id": request.task_id,
            "baseline": request.baseline.model_dump(),
            "result": request.result.model_dump(),
            "evaluation": request.evaluation.model_dump(),
        },
    )
    _write_json(target / "execution_record.json", request.execution_record)
    names = [*coverage, f"{prefix}/evaluation.json", f"{prefix}/execution_record.json"]
    artifacts = [
        *request.result.artifacts,
        *request.derived_artifacts,
        *(describe_artifact(directory, name) for name in names),
    ]
    _write_json(
        target / "artifact-index.json",
        {
            "task_id": request.task_id,
            "baseline_version": request.baseline.version,
            "report_fingerprint": render_key,
            "reports": coverage,
            "artifacts": [item.model_dump() for item in artifacts],
        },
    )
    artifacts.append(describe_artifact(directory, f"{prefix}/artifact-index.json"))
    for item in artifacts:
        verify_artifact(directory, item)
    return artifacts


def finalize_execution_record(
    request: ReportRequest, directory: Path, artifacts: list[Artifact]
) -> list[Artifact]:
    prefix = f"delivery-v{request.baseline.version}"
    execution_name = f"{prefix}/execution_record.json"
    index_name = f"{prefix}/artifact-index.json"
    index_artifact = next(item for item in artifacts if item.path == index_name)
    index = json.loads(verify_artifact(directory, index_artifact).read_text("utf-8"))
    _write_json(directory / execution_name, request.execution_record)
    updated = [
        describe_artifact(directory, execution_name) if item.path == execution_name else item
        for item in artifacts
        if item.path != index_name
    ]
    index["artifacts"] = [item.model_dump() for item in updated]
    _write_json(directory / index_name, index)
    updated.append(describe_artifact(directory, index_name))
    for item in updated:
        verify_artifact(directory, item)
    return updated
