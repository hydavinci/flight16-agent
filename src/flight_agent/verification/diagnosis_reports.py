from pathlib import Path

from docx import Document

from ..runtime.evidence import load_run
from ..session import _write_json
from .artifacts import describe_artifact, verify_artifact
from .contracts import Artifact, ConversationMessage, DiagnosisTask
from .diagnosis_contracts import validate_assessments, validate_history
from .diagnostic_data import DiagnosticOutput
from .intake import safe_directory


def deliver_diagnosis_report(
    task: DiagnosisTask, messages: list[ConversationMessage]
) -> list[Artifact]:
    if not task.baseline.confirmed or task.termination is None:
        raise ValueError(
            "Diagnosis delivery requires a confirmed baseline and recorded stop reason"
        )
    current = validate_history(
        task.baseline.hypotheses, task.plan_revisions, task.rounds, task.analysis_runs
    )
    validate_assessments(current, task.assessments, task.analysis_runs)
    if task.rounds and task.assessments != task.rounds[-1].review.assessments:
        raise ValueError("Recorded assessments differ from the final evidence review")
    if task.termination.reason == "evidence_sufficient" and (
        not task.rounds or task.rounds[-1].review.stop_reason != "evidence_sufficient"
        or task.assessments != task.rounds[-1].review.assessments
    ):
        raise ValueError("A sufficient-evidence report requires the recorded final assessment")
    summary = task.rounds[-1].review.summary if task.rounds else None
    revisions = [
        {"plan_version": item.plan_version, "revision": item.review.hypothesis_revision.model_dump(
            mode="json"
        )}
        for item in task.rounds if item.review.hypothesis_revision
    ]
    directory = Path(task.result_directory)
    safe_directory(directory)
    prefix = f"delivery-v{task.baseline.version}"
    target = directory / prefix
    safe_directory(target)
    target.mkdir(exist_ok=True)
    for name in (
        "diagnosis.docx", "diagnosis.tmp.docx", "diagnosis.json",
        "execution_record.json", "artifact-index.json",
    ):
        if (target / name).is_symlink():
            raise ValueError("Unsafe diagnostic delivery path")
    for artifact in task.artifacts:
        verify_artifact(directory, artifact)
    for run in task.analysis_runs:
        if load_run(directory, run.id, run.spec) != run:
            raise ValueError("Diagnosis report evidence does not match its immutable receipt")
    document = Document()
    document.add_heading(task.proposal.title, 0)
    document.add_paragraph(
        f"D2 evidence assessment; task {task.id}; baseline v{task.baseline.version}."
    )
    document.add_paragraph(
        "Hypothesis support/refutation is not proof of a unique physical cause, "
        "flight qualification, or engineering acceptance. Real labeled data and "
        "independent references are required."
    )
    if task.project.manifest.synthetic:
        document.add_paragraph("SYNTHETIC DATA ONLY: this is not flight-test or MATLAB validation.")
    document.add_heading("Confirmed scope and versions", 1)
    document.add_paragraph(f"Project fingerprint: {task.baseline.project_fingerprint}")
    document.add_paragraph("Design references: " + task.baseline.design.model_dump_json())
    document.add_paragraph("Allowed capabilities: " + ", ".join(task.baseline.allowed_capabilities))
    document.add_paragraph("Resource budget: " + task.baseline.budget.model_dump_json())
    document.add_paragraph("Conditions: " + str(task.project.manifest.conditions))
    document.add_heading("Confirmed initial hypotheses", 2)
    for hypothesis in task.baseline.hypotheses:
        document.add_paragraph(
            f"{hypothesis.id} [{hypothesis.priority.upper()}]: {hypothesis.statement}"
        )
    document.add_heading("Problem localization", 1)
    if summary is None:
        document.add_paragraph(
            "Structured diagnosis unavailable: no structured summary was recorded. "
            "Do not infer problem localization or model/control responsibility."
        )
    else:
        document.add_paragraph("Problem: " + summary.problem)
        document.add_paragraph("Localization: " + summary.localization)
        document.add_paragraph("Model/control responsibility: " + summary.responsibility)
        document.add_paragraph("Attribution limits: " + summary.reason)
        document.add_paragraph("Related hypotheses: " + ", ".join(summary.hypothesis_ids))
        for citation in summary.evidence:
            document.add_paragraph(
                f"Evidence: run {citation.run_id}"
                + (f", metric {citation.metric}" if citation.metric else "")
            )
        document.add_heading("Recommendations", 2)
        for recommendation in summary.recommendations:
            document.add_paragraph(recommendation)
    document.add_heading("Current hypotheses and evidence", 1)
    hypotheses = {item.id: item for item in current}
    for assessment in task.assessments:
        hypothesis = hypotheses[assessment.hypothesis_id]
        document.add_heading(
            f"{assessment.hypothesis_id}: {assessment.status.upper()} "
            f"/ {hypothesis.priority.upper()}", 2
        )
        document.add_paragraph(hypothesis.statement)
        document.add_paragraph(assessment.reason)
        for citation in assessment.evidence:
            document.add_paragraph(
                f"Evidence: run {citation.run_id}"
                + (f", metric {citation.metric}" if citation.metric else "")
            )
    document.add_heading("Append-only analysis evidence", 1)
    for run in task.analysis_runs:
        output = DiagnosticOutput.model_validate(run.output)
        document.add_heading(f"{run.spec.capability} / {run.id}", 2)
        document.add_paragraph(f"Full input fingerprint: {run.spec.fingerprint}")
        document.add_paragraph("Inputs: " + str(run.spec.inputs))
        for name, metric in output.metrics.items():
            document.add_paragraph(f"{name}: {metric.value:g} {metric.unit}")
        for observation in output.observations:
            document.add_paragraph("Observation: " + observation)
        for unknown in output.unknowns:
            document.add_paragraph("Unknown: " + unknown)
    document.add_heading("Adaptive plan revisions", 1)
    for plan in task.plan_revisions:
        document.add_paragraph(
            f"Plan {plan.version}: {plan.reason}; previous evidence: {', '.join(plan.evidence)}"
        )
    document.add_heading("Hypothesis revision history", 1)
    if not revisions:
        document.add_paragraph("No hypothesis revisions were recorded.")
    for round_ in task.rounds:
        if revision := round_.review.hypothesis_revision:
            document.add_heading(f"After round {round_.plan_version}", 2)
            document.add_paragraph(revision.reason)
            for hypothesis in revision.hypotheses:
                document.add_paragraph(
                    f"{hypothesis.id} [{hypothesis.priority.upper()}]: {hypothesis.statement}"
                )
            for citation in revision.evidence:
                document.add_paragraph(
                    f"Evidence: run {citation.run_id}"
                    + (f", metric {citation.metric}" if citation.metric else "")
                )
    document.add_heading("Termination and remaining unknowns", 1)
    document.add_paragraph(task.termination.reason + ": " + task.termination.detail)
    for unknown in task.termination.unknowns:
        document.add_paragraph(unknown)
    temporary = target / "diagnosis.tmp.docx"
    document.save(str(temporary))
    temporary.replace(target / "diagnosis.docx")
    payload = task.model_dump(mode="json")
    payload["stage"] = "completed"
    findings = {
        "current_hypotheses": [item.model_dump(mode="json") for item in current],
        "hypothesis_revisions": revisions,
        "summary": summary.model_dump(mode="json") if summary else None,
    }
    _write_json(
        target / "diagnosis.json",
        {
            "task_id": task.id,
            "baseline": task.baseline.model_dump(mode="json"),
            **findings,
            "assessments": [item.model_dump(mode="json") for item in task.assessments],
            "termination": task.termination.model_dump(mode="json"),
        },
    )
    _write_json(
        target / "execution_record.json",
        {
            **payload, **findings, "status": "completed",
            "messages": [item.model_dump() for item in messages],
        },
    )
    names = ["diagnosis.docx", "diagnosis.json", "execution_record.json"]
    artifacts = [
        *(item for item in task.artifacts if not item.path.startswith(prefix + "/")),
        *(describe_artifact(directory, f"{prefix}/{name}") for name in names),
    ]
    _write_json(
        target / "artifact-index.json",
        {
            "task_id": task.id, "baseline_version": task.baseline.version,
            "reports": [f"{prefix}/diagnosis.docx"],
            "artifacts": [item.model_dump() for item in artifacts],
        },
    )
    artifacts.append(describe_artifact(directory, f"{prefix}/artifact-index.json"))
    for artifact in artifacts:
        verify_artifact(directory, artifact)
    return artifacts
